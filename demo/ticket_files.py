"""測試工單檔：依主題類別再依來源分資料夾，json／csv／xlsx 讀成同一個 `Ticket`。

目錄契約：

```text
demo/test-tickets/{category}/{email|discord|github_issue}/{json|csv|xlsx}/tickets.<ext>
```

第一層是主題資料夾（例如 `prepare-meeting`），名稱可自訂，不必先登記。
`source` 必須與來源資料夾名稱一致。本模組不碰 DynamoDB、不呼叫 Lambda。
xlsx 只在本機 dev 依賴 `openpyxl` 裡解析，不進 Lambda runtime。
"""

from __future__ import annotations

import csv
import json
from collections.abc import Mapping, Sequence
from pathlib import Path
from typing import Any

from pydantic import ValidationError

from training_kb.clock import to_iso
from training_kb.errors import ContentError
from training_kb.models import Ticket, TicketSource

SOURCES: tuple[str, ...] = (
    str(TicketSource.EMAIL), str(TicketSource.DISCORD), str(TicketSource.GITHUB_ISSUE))
"""資料夾名稱；順序是填表的人較常見的 email → discord → github_issue。"""

FORMATS: tuple[str, ...] = ("json", "csv", "xlsx")
COLUMNS: tuple[str, ...] = (
    "id", "source", "text", "author", "ts", "project_id", "cluster_id", "feature_ids",
    "embedding")
DEFAULT_TICKETS_DIR = "demo/test-tickets"
SEED_CATEGORIES: tuple[str, ...] = (
    "prepare-meeting", "share-summary", "notification-settings", "weekly-digest")
"""匯出腳本預先建好的主題資料夾；`load_all` 仍會掃第一層任何非來源名的目錄。"""
NOTICE = ("合成資料示範（SYNTHETIC）：測試上傳用工單，不是真實觀測、不含任何真人資料、帳號或金鑰。")
GITHUB_ISSUE_HINT = (
    "github_issue 不能走 training-kb-import 的精簡 Ticket 檔；"
    "正式入口是 GitHub webhook Function URL。")


def ticket_row(ticket: Ticket) -> dict[str, Any]:
    """Ticket → 三種檔案共用的一列；時間固定 Z、list 原樣留下。"""
    return {
        "id": ticket.id, "source": str(ticket.source), "text": ticket.text,
        "author": ticket.author, "ts": to_iso(ticket.ts), "project_id": ticket.project_id,
        "cluster_id": ticket.cluster_id, "feature_ids": list(ticket.feature_ids),
        "embedding": ticket.embedding,
    }


def parse_ticket(row: Mapping[str, Any], *, name: str, index: int) -> Ticket:
    try:
        return Ticket.model_validate(row)
    except ValidationError as error:
        raise ContentError(
            f"{name} 第 {index + 1} 筆不符 Ticket：{error.error_count()} 個欄位"
        ) from error


def _empty(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _parse_list(value: object, *, field: str, name: str, index: int) -> list[Any] | None:
    """CSV／xlsx 儲存格 → list；空字串是「沒有」。embedding 空回 None，feature_ids 空回 []。"""
    if _empty(value):
        return None if field == "embedding" else []
    if isinstance(value, list):
        return value
    text = str(value).strip()
    try:
        parsed: Any = json.loads(text)
    except json.JSONDecodeError as error:
        raise ContentError(
            f"{name} 第 {index + 1} 筆的 {field} 不是 JSON 陣列"
        ) from error
    if parsed is None:
        return None if field == "embedding" else []
    if not isinstance(parsed, list):
        raise ContentError(f"{name} 第 {index + 1} 筆的 {field} 必須是陣列")
    return parsed


def _cell_row(raw: Mapping[str, Any], *, name: str, index: int) -> dict[str, Any]:
    row = {key: raw.get(key) for key in COLUMNS}
    cluster = row["cluster_id"]
    row["cluster_id"] = None if _empty(cluster) else cluster
    row["feature_ids"] = _parse_list(
        row["feature_ids"], field="feature_ids", name=name, index=index) or []
    row["embedding"] = _parse_list(
        row["embedding"], field="embedding", name=name, index=index)
    return row


def _format_of(path: Path) -> str:
    suffix = path.suffix.lower().lstrip(".")
    if suffix not in FORMATS:
        raise ContentError(f"不支援的工單檔格式：{path}（只要 json／csv／xlsx）")
    return suffix


def infer_source(path: Path) -> str | None:
    """`.../{source}/{format}/file` 才認得出來源資料夾；否則回 None，改看 Ticket.source。

    主題層不參與判斷：`prepare-meeting/email/json/tickets.json` 與舊的
    `email/json/tickets.json` 都認成 `email`。
    """
    parts = path.resolve().parts
    if len(parts) < 3:
        return None
    source, fmt = parts[-3], parts[-2]
    if source in SOURCES and fmt in FORMATS:
        return source
    return None


def list_categories(directory: Path) -> tuple[str, ...]:
    """第一層不是來源名、也不是隱藏目錄的資料夾，就是主題類別。"""
    root = Path(directory)
    if not root.is_dir():
        return ()
    return tuple(sorted(
        child.name for child in root.iterdir()
        if child.is_dir() and child.name not in SOURCES and not child.name.startswith(".")))


def _reserved_category(category: str) -> None:
    if category in SOURCES or category in FORMATS or category.startswith("."):
        raise ContentError(
            f"類別資料夾名稱不能是來源、格式或隱藏目錄：{category}")


def _check_source(tickets: Sequence[Ticket], *, expected: str, name: str) -> None:
    for index, ticket in enumerate(tickets):
        if str(ticket.source) != expected:
            raise ContentError(
                f"{name} 第 {index + 1} 筆 source={ticket.source}，"
                f"與資料夾 {expected} 不一致")


def load_json(path: Path) -> tuple[Ticket, ...]:
    try:
        payload: Any = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as error:
        raise ContentError(f"不是合法 JSON：{path}：{error}") from error
    if not isinstance(payload, dict):
        raise ContentError(f"工單 JSON 最外層必須是物件：{path}")
    items = payload.get("items")
    if not isinstance(items, list):
        raise ContentError(f"工單 JSON 缺少 items 陣列：{path}")
    tickets: list[Ticket] = []
    for index, row in enumerate(items):
        if not isinstance(row, Mapping):
            raise ContentError(
                f"{path.name} 第 {index + 1} 筆必須是物件，實際是 {type(row).__name__}")
        tickets.append(parse_ticket(row, name=path.name, index=index))
    return tuple(tickets)


def load_csv(path: Path) -> tuple[Ticket, ...]:
    with path.open(encoding="utf-8", newline="") as handle:
        reader = csv.DictReader(handle)
        if reader.fieldnames is None:
            raise ContentError(f"CSV 缺少欄名：{path}")
        missing = [column for column in COLUMNS if column not in reader.fieldnames]
        if missing:
            raise ContentError(f"CSV 缺少欄位 {missing}：{path}")
        tickets: list[Ticket] = []
        for index, raw in enumerate(reader):
            tickets.append(parse_ticket(
                _cell_row(raw, name=path.name, index=index), name=path.name, index=index))
    return tuple(tickets)


def _openpyxl() -> Any:
    try:
        import openpyxl  # type: ignore[import-untyped]
    except ImportError as error:
        raise ContentError(
            "讀寫 xlsx 需要 openpyxl；請在本機安裝 dev 依賴（不會進 Lambda）"
        ) from error
    return openpyxl


def load_xlsx(path: Path) -> tuple[Ticket, ...]:
    workbook = _openpyxl().load_workbook(path, data_only=True)
    sheet = workbook.active
    if sheet is None:
        raise ContentError(f"xlsx 沒有工作表：{path}")
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise ContentError(f"xlsx 是空的：{path}")
    header = [str(cell).strip() if cell is not None else "" for cell in rows[0]]
    missing = [column for column in COLUMNS if column not in header]
    if missing:
        raise ContentError(f"xlsx 缺少欄位 {missing}：{path}")
    tickets: list[Ticket] = []
    for index, values in enumerate(rows[1:]):
        raw = {key: (values[header.index(key)] if header.index(key) < len(values) else None)
               for key in COLUMNS}
        tickets.append(parse_ticket(
            _cell_row(raw, name=path.name, index=index), name=path.name, index=index))
    return tuple(tickets)


_LOADERS = {"json": load_json, "csv": load_csv, "xlsx": load_xlsx}


def load_file(path: Path, *, expected_source: str | None = None) -> tuple[Ticket, ...]:
    """讀一個工單檔。路徑落在來源資料夾時，每一筆 source 都要對得上。"""
    fmt = _format_of(path)
    if not path.is_file():
        raise ContentError(f"找不到工單檔：{path}")
    tickets = _LOADERS[fmt](path)
    source = expected_source if expected_source is not None else infer_source(path)
    if source is not None:
        _check_source(tickets, expected=source, name=str(path))
    return tickets


def slot_dir(directory: Path, source: str, fmt: str, *, category: str | None = None) -> Path:
    if source not in SOURCES:
        raise ContentError(f"未知來源資料夾：{source}；可用的是 {list(SOURCES)}")
    if fmt not in FORMATS:
        raise ContentError(f"未知格式：{fmt}；可用的是 {list(FORMATS)}")
    if category is None:
        return directory / source / fmt
    _reserved_category(category)
    return directory / category / source / fmt


def _slot_folders(directory: Path, source: str, fmt: str,
                  *, category: str | None) -> list[Path]:
    """有主題資料夾就掃每一類；沒有就退回舊的 `{source}/{format}`。"""
    root = Path(directory)
    if category is not None:
        return [slot_dir(root, source, fmt, category=category)]
    names = list_categories(root)
    if names:
        return [slot_dir(root, source, fmt, category=name) for name in names]
    return [slot_dir(root, source, fmt)]


def load_slot(directory: Path, source: str, fmt: str,
              *, category: str | None = None) -> tuple[Ticket, ...]:
    candidates = _slot_folders(directory, source, fmt, category=category)
    folders = [folder for folder in candidates if folder.is_dir()]
    if not folders:
        raise ContentError(f"缺少測試工單目錄：{candidates[0]}")
    tickets: list[Ticket] = []
    seen_files = False
    for folder in folders:
        paths = sorted(path for path in folder.iterdir()
                       if path.is_file() and path.suffix.lower() == f".{fmt}")
        if not paths:
            if category is not None:
                raise ContentError(f"{folder} 裡沒有 .{fmt} 檔")
            continue
        seen_files = True
        for path in paths:
            tickets.extend(load_file(path, expected_source=source))
    if not seen_files:
        raise ContentError(f"{folders[0]} 裡沒有 .{fmt} 檔")
    return tuple(tickets)


def load_all(directory: Path) -> tuple[Ticket, ...]:
    """掃完整棵樹；同一 id 出現多次時內容必須相同，只留一筆。"""
    root = Path(directory)
    if not root.is_dir():
        raise ContentError(f"缺少測試工單目錄：{root}")
    by_id: dict[str, Ticket] = {}
    bases = [root / name for name in list_categories(root)] or [root]
    for base in bases:
        for source in SOURCES:
            for fmt in FORMATS:
                folder = base / source / fmt
                if not folder.is_dir():
                    continue
                for path in sorted(folder.iterdir()):
                    if not path.is_file() or path.suffix.lower() != f".{fmt}":
                        continue
                    for ticket in load_file(path, expected_source=source):
                        existing = by_id.get(ticket.id)
                        if existing is None:
                            by_id[ticket.id] = ticket
                            continue
                        if ticket_row(existing) != ticket_row(ticket):
                            raise ContentError(
                                f"工單 {ticket.id} 在測試目錄裡出現不一致的內容")
    return tuple(by_id[key] for key in sorted(by_id))


def find_ticket(directory: Path, ticket_id: str) -> Ticket | None:
    return next((row for row in load_all(directory) if row.id == ticket_id), None)


def dump_json(path: Path, tickets: Sequence[Ticket]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = {"_notice": NOTICE, "synthetic": True,
               "items": [ticket_row(row) for row in tickets]}
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def _csv_cell(column: str, row: Mapping[str, Any]) -> str:
    value = row[column]
    if column in {"feature_ids", "embedding"}:
        return "" if value in (None, []) else json.dumps(value, ensure_ascii=False)
    if value is None:
        return ""
    return str(value)


def dump_csv(path: Path, tickets: Sequence[Ticket]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(COLUMNS), extrasaction="ignore")
        writer.writeheader()
        for ticket in tickets:
            row = ticket_row(ticket)
            writer.writerow({column: _csv_cell(column, row) for column in COLUMNS})


def dump_xlsx(path: Path, tickets: Sequence[Ticket]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = _openpyxl().Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(list(COLUMNS))
    for ticket in tickets:
        row = ticket_row(ticket)
        sheet.append([_csv_cell(column, row) for column in COLUMNS])
    workbook.save(path)


def dump_source(directory: Path, source: str, tickets: Sequence[Ticket],
                *, category: str | None = None) -> None:
    """同一個來源寫出 json／csv／xlsx 各一份 `tickets.<ext>`。"""
    dump_json(slot_dir(directory, source, "json", category=category) / "tickets.json", tickets)
    dump_csv(slot_dir(directory, source, "csv", category=category) / "tickets.csv", tickets)
    dump_xlsx(slot_dir(directory, source, "xlsx", category=category) / "tickets.xlsx", tickets)
